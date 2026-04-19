from __future__ import annotations

import csv
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
FIGURES = REPORT / "figures"
TABLES = REPORT / "tables"
OUTPUT_DIR = ROOT / "output" / "presentation"
OUTPUT_FILE = OUTPUT_DIR / "BRSM_Project_Phase1_Presentation.pptx"


BG = RGBColor(245, 247, 250)
NAVY = RGBColor(20, 33, 61)
BLUE = RGBColor(51, 102, 153)
TEAL = RGBColor(35, 132, 136)
ORANGE = RGBColor(203, 108, 55)
TEXT = RGBColor(40, 40, 40)
MUTED = RGBColor(92, 102, 112)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(228, 233, 239)
GREEN = RGBColor(45, 125, 80)


def load_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def fmt3(value: str) -> str:
    return f"{float(value):.3f}"


def add_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.42)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()


def add_title(slide, title: str, subtitle: str | None = None, owner: str | None = None):
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.55), Inches(8.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.58), Inches(1.12), Inches(8.6), Inches(0.45))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        run.font.name = "Aptos"
        run.font.size = Pt(11.5)
        run.font.color.rgb = MUTED

    if owner:
        owner_box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.75), Inches(0.62), Inches(1.9), Inches(0.42)
        )
        owner_box.fill.solid()
        owner_box.fill.fore_color.rgb = WHITE
        owner_box.line.color.rgb = LIGHT
        tf = owner_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"Owner: {owner}"
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = BLUE


def add_bullets(slide, bullets: list[str], left, top, width, height, font_size=18, color=TEXT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
        p.bullet = True
    return box


def add_note(slide, note_text: str):
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.text = note_text


def fit_image(slide, image_path: Path, left, top, width, height):
    with Image.open(image_path) as img:
        img_w, img_h = img.size
    img_ratio = img_w / img_h
    box_ratio = width / height

    if img_ratio > box_ratio:
        final_w = width
        final_h = width / img_ratio
        final_left = left
        final_top = top + (height - final_h) / 2
    else:
        final_h = height
        final_w = height * img_ratio
        final_top = top
        final_left = left + (width - final_w) / 2
    slide.shapes.add_picture(str(image_path), final_left, final_top, width=final_w, height=final_h)


def add_metric_card(slide, x, y, w, h, title, lines, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LIGHT

    title_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.3))
    p = title_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = accent

    text_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.42), w - Inches(0.36), h - Inches(0.52))
    tf = text_box.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT
        p.space_after = Pt(5)


def add_pipeline_box(slide, x, y, w, h, title, body, accent):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = accent
    box.line.width = Pt(1.6)

    tbox = slide.shapes.add_textbox(x + Inches(0.16), y + Inches(0.12), w - Inches(0.3), Inches(0.3))
    p = tbox.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Aptos"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = accent

    bbox = slide.shapes.add_textbox(x + Inches(0.16), y + Inches(0.45), w - Inches(0.3), h - Inches(0.58))
    tf = bbox.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(body):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(11.2)
        p.font.color.rgb = TEXT
        p.space_after = Pt(4)


def add_arrow(slide, x, y, w, h):
    arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, x, y, w, h)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = LIGHT
    arrow.line.fill.background()


def build_presentation():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    sample_rows = load_csv_rows(TABLES / "sample_overview.csv")
    encoding_rows = load_csv_rows(TABLES / "encoding_descriptives.csv")
    memory_rows = load_csv_rows(TABLES / "memory_descriptives.csv")
    within_rows = load_csv_rows(TABLES / "within_condition_tests.csv")

    sample_map = {row["condition"]: row for row in sample_rows}
    encoding_map = {(r["condition"], r["boundary_position"]): r for r in encoding_rows}
    memory_map = {(r["condition"], r["metric"], r["boundary_position"]): r for r in memory_rows}

    def row_for(condition: str, metric: str):
        return next(r for r in within_rows if r["analysis_family"] == "within_condition_anova" and r["condition"] == condition and r["metric"] == metric)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.2), Inches(11.9), Inches(4.9))
    hero.fill.solid()
    hero.fill.fore_color.rgb = WHITE
    hero.line.color.rgb = LIGHT
    add_title(slide, "Behavioural Research in Statistical Methods", "Project Report 1 | Mnemonic Similarity Task (MST)", owner="Shared")
    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(10.8), Inches(2.0))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Event Boundaries and Recognition Memory"
    r.font.name = "Aptos Display"
    r.font.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = NAVY
    p = tf.add_paragraph()
    p.text = "Team Earphones"
    p.font.name = "Aptos"
    p.font.size = Pt(18)
    p.font.color.rgb = BLUE
    p = tf.add_paragraph()
    p.text = "Sambu Aneesh | Renu Sree Vyshnavi | Pavan Harshit Chinta"
    p.font.name = "Aptos"
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT
    add_bullets(
        slide,
        [
            "Exploratory analysis of MST data to test how memory changes around event boundaries.",
            "Focus on encoding response time, recognition memory (REC), and lure discrimination index (LDI).",
            "Presentation flow: context -> methods -> results -> interpretation -> conclusion.",
        ],
        Inches(1.0), Inches(4.0), Inches(11.0), Inches(1.7), font_size=16,
    )
    add_note(
        slide,
        "Open by introducing the project title, the team, and the central question. "
        "Explain that this is Phase 1 of the project and the goal here is exploratory rather than a final causal claim. "
        "Say that the deck is structured in the same order as the report: first the experiment context, then preprocessing and analysis choices, then the main results and what they mean for Phase 2."
    )

    # Slide 2
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Background and Study Motivation", "Why event boundaries may matter for memory", owner="Sambu")
    add_bullets(
        slide,
        [
            "Event segmentation theory suggests people naturally divide continuous experience into meaningful events.",
            "Boundaries may change how information is encoded, updated, and later retrieved from memory.",
            "The Mnemonic Similarity Task (MST) is useful because it goes beyond simple recognition and tests fine-grained memory discrimination.",
            "The target paper reported two directional effects: weaker post-boundary recognition and stronger pre-boundary lure discrimination.",
        ],
        Inches(0.75), Inches(1.55), Inches(6.0), Inches(4.8), font_size=18,
    )
    add_metric_card(
        slide,
        Inches(7.1), Inches(1.75), Inches(5.35), Inches(3.9),
        "Why this study matters",
        [
            "Links cognitive theory with behavioral data.",
            "Tests whether boundary cues only affect processing speed or also affect memory quality.",
            "Checks whether the published pattern replicates in our dataset.",
        ],
        accent=TEAL,
    )
    add_note(
        slide,
        "Explain the conceptual bridge between event segmentation and MST. "
        "The key motivation is that boundaries are not only perceptual transitions; they may reorganize memory. "
        "Mention that the target paper motivates the expected pattern: a cost immediately after a boundary and a possible benefit just before it. "
        "Keep the emphasis on why the experiment is interesting, not yet on our own findings."
    )

    # Slide 3
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Research Questions and Experimental Context", "How the experiment was structured", owner="Sambu")
    add_bullets(
        slide,
        [
            "RQ1: Do event boundaries slow encoding response time?",
            "RQ2: Is recognition memory weaker for post-boundary items?",
            "RQ3: Do pre-boundary items show better lure discrimination?",
            "RQ4: Do these patterns differ across the three boundary manipulations?",
            "RQ5: Do lure bins show the expected MST trend as lures become less similar?",
        ],
        Inches(0.75), Inches(1.55), Inches(5.9), Inches(3.2), font_size=18,
    )
    add_metric_card(
        slide,
        Inches(6.95), Inches(1.55), Inches(5.55), Inches(1.45),
        "Three task conditions",
        [
            "Item Shift Only (`item_only`)",
            "Item + Task Shift (`both`)",
            "Task Shift Only (`task_only`)",
        ],
        accent=BLUE,
    )
    add_metric_card(
        slide,
        Inches(6.95), Inches(3.15), Inches(5.55), Inches(2.15),
        "Boundary positions inside each 7-item event",
        [
            "`post`: first item after a boundary",
            "`mid`: five middle items",
            "`pre`: last item before the next boundary",
        ],
        accent=ORANGE,
    )
    add_note(
        slide,
        "Walk through the five research questions clearly. "
        "Then describe the experiment setup: there are three versions of the task, which differ in the type of boundary cue. "
        "Also define post, mid, and pre positions because that distinction is used throughout the rest of the analysis. "
        "This slide should make later results easy to follow."
    )

    # Slide 4
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Dataset and Conditions", "Cleaned participant sample and task structure", owner="Sambu")
    add_metric_card(
        slide, Inches(0.85), Inches(1.6), Inches(3.85), Inches(1.5),
        "Item Shift Only", [f"Paired participants: {sample_map['item_only']['participants']}", "Encoding trials per participant: 280", "Test trials per participant: 150"], accent=BLUE
    )
    add_metric_card(
        slide, Inches(4.75), Inches(1.6), Inches(3.85), Inches(1.5),
        "Item + Task Shift", [f"Paired participants: {sample_map['both']['participants']}", "Encoding trials per participant: 280", "Test trials per participant: 150"], accent=TEAL
    )
    add_metric_card(
        slide, Inches(8.65), Inches(1.6), Inches(3.85), Inches(1.5),
        "Task Shift Only", [f"Paired participants: {sample_map['task_only']['participants']}", "Encoding trials per participant: 280", "Test trials per participant: 150"], accent=ORANGE
    )
    fit_image(slide, FIGURES / "phase1_dataset_overview.png", Inches(0.85), Inches(3.35), Inches(7.0), Inches(3.45))
    add_bullets(
        slide,
        [
            "Only complete task-test pairs were retained for analysis.",
            "Participants with missing task or test files were excluded.",
            "The paired sample sizes are balanced enough for condition-wise repeated-measures analysis.",
        ],
        Inches(8.2), Inches(3.65), Inches(4.2), Inches(2.1), font_size=16,
    )
    add_note(
        slide,
        "Use this slide to show that the dataset was organized before any statistical testing. "
        "State the paired sample sizes in each condition and remind the audience that incomplete sessions were excluded. "
        "Point to the dataset overview figure and note that file resolution and participant pairing were a real preprocessing step, not assumed clean input."
    )

    # Slide 5
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Data Preparation and Preprocessing", "How raw task and test files became analysis-ready data", owner="Pavan")
    x_positions = [Inches(0.6), Inches(3.35), Inches(6.1), Inches(8.85), Inches(11.1)]
    widths = [Inches(2.35)] * 5
    bodies = [
        ["Read raw CSV files from all three MST condition folders.", "Matched task and test files participant by participant."],
        ["Used clean pairs directly.", "For duplicates, paired the latest test file with the previous task file."],
        ["Dropped incomplete sessions with only task or only test data.", "Kept participants as the analytical unit."],
        ["Retained 280 non-practice encoding trials.", "Derived `post`, `mid`, and `pre` from 7-item event structure."],
        ["Retained 150 non-empty test trials.", "Parsed filenames to classify target, lure, and foil items."],
    ]
    titles = ["Raw Input", "File Pairing", "Exclusions", "Encoding Setup", "Test Setup"]
    accents = [BLUE, TEAL, ORANGE, BLUE, TEAL]
    for x, w, title, body, accent in zip(x_positions, widths, titles, bodies, accents):
        add_pipeline_box(slide, x, Inches(2.0), w, Inches(2.35), title, body, accent)
    for x in [Inches(2.95), Inches(5.7), Inches(8.45), Inches(10.7)]:
        add_arrow(slide, x, Inches(2.7), Inches(0.28), Inches(0.4))
    add_bullets(
        slide,
        [
            "Response time was derived from in-trial response when available, or from post-stimulus RT adjusted by +3 seconds when the response came after image offset.",
            "This preprocessing ensures later condition comparisons are based on clean, comparable participant summaries rather than raw trial counts.",
        ],
        Inches(0.85), Inches(5.0), Inches(11.9), Inches(1.25), font_size=15,
    )
    add_note(
        slide,
        "Present this as the analysis pipeline. "
        "Explain the file-pairing decisions carefully because they justify why the sample sizes are what they are. "
        "Mention the duplicate-resolution rule, the exclusion of incomplete sessions, and the boundary-position labeling logic. "
        "End by saying that the participant, not the individual trial, is the analytical unit."
    )

    # Slide 6
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Variables and Scoring Metrics", "What was measured and how each score was defined", owner="Pavan")
    add_metric_card(
        slide, Inches(0.8), Inches(1.7), Inches(3.8), Inches(2.2),
        "Independent variables",
        ["Condition: `item_only`, `both`, `task_only`", "Boundary position: `post`, `mid`, `pre`"], accent=BLUE
    )
    add_metric_card(
        slide, Inches(4.8), Inches(1.7), Inches(3.8), Inches(2.2),
        "Encoding variable",
        ["Mean encoding response time", "Response completeness by boundary position"], accent=TEAL
    )
    add_metric_card(
        slide, Inches(8.8), Inches(1.7), Inches(3.8), Inches(2.2),
        "Memory variables",
        ["REC: bias-corrected target recognition", "LDI: bias-corrected lure discrimination"], accent=ORANGE
    )
    formula_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.35), Inches(11.25), Inches(1.55))
    formula_box.fill.solid()
    formula_box.fill.fore_color.rgb = WHITE
    formula_box.line.color.rgb = LIGHT
    ftxt = slide.shapes.add_textbox(Inches(1.3), Inches(4.7), Inches(10.8), Inches(0.8))
    p = ftxt.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = 'REC = P("old" | Target) - P("old" | Foil)      |      LDI = P("similar" | Lure) - P("similar" | Foil)'
    r.font.name = "Aptos"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = NAVY
    add_bullets(
        slide,
        [
            "A sensitivity check also computed an alternate LDI-style measure by combining `new` and `similar` as lure rejection responses.",
            "Lure-similarity bins were attached to each image to test whether discrimination improves as lures become less similar.",
        ],
        Inches(1.0), Inches(6.15), Inches(11.1), Inches(0.8), font_size=14,
    )
    add_note(
        slide,
        "Define the variables slowly because this is where the audience needs clarity. "
        "Explain REC as corrected recognition and LDI as corrected lure discrimination. "
        "Make it clear that both metrics subtract foil-based response tendencies to reduce response bias. "
        "Then mention the alternate scoring sensitivity check and the lure-bin analysis as supporting analyses."
    )

    # Slide 7
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Statistical Tests and Analysis Plan", "Hypotheses and inferential strategy", owner="Pavan")
    add_metric_card(
        slide, Inches(0.8), Inches(1.65), Inches(5.7), Inches(3.95),
        "Hypotheses",
        [
            "H1 (Encoding RT): post-boundary items will be slower, especially when the task rule changes.",
            "H1 (Recognition): post-boundary REC will be lower than mid-event REC, strongest in `both`.",
            "H1 (Lure discrimination): pre-boundary LDI may be higher than post or mid, following the target paper.",
        ],
        accent=BLUE,
    )
    add_metric_card(
        slide, Inches(6.8), Inches(1.65), Inches(5.7), Inches(3.95),
        "Analysis workflow",
        [
            "1. Compute participant-level summaries for each condition and boundary position.",
            "2. Visualize means and individual structure before testing.",
            "3. Run repeated-measures ANOVA within each condition.",
            "4. Follow informative ANOVAs with Holm-corrected paired t-tests.",
            "5. Compare contrast scores across conditions when needed.",
        ],
        accent=TEAL,
    )
    add_bullets(
        slide,
        [
            "This strategy keeps the analysis interpretable and aligned with the course emphasis on subject-level inference and multiple-comparison control.",
        ],
        Inches(0.95), Inches(6.05), Inches(11.3), Inches(0.55), font_size=14,
    )
    add_note(
        slide,
        "This is your analysis slide, so emphasize why the tests match the design. "
        "Because boundary position varies within participants, repeated-measures ANOVA is appropriate within each condition. "
        "Then explain that Holm correction was used for the pairwise follow-up tests to control false positives. "
        "Close by linking the hypotheses directly to the three main dependent variables."
    )

    both_rt = row_for("both", "encoding_rt")
    task_rt = row_for("task_only", "encoding_rt")
    item_rt = row_for("item_only", "encoding_rt")

    # Slide 8
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Results: Data Quality and Response Profiles", "Participants used the expected response categories", owner="Renu")
    fit_image(slide, FIGURES / "phase1_response_profiles.png", Inches(0.65), Inches(1.55), Inches(8.1), Inches(5.5))
    add_bullets(
        slide,
        [
            "Targets were mostly labeled `old`, lures were mostly labeled `similar`, and foils were mostly labeled `new`.",
            "This diagonal pattern is a basic but important quality-control check before interpreting memory indices.",
            "It supports that participants understood the response options and were not responding randomly.",
        ],
        Inches(8.95), Inches(2.0), Inches(3.7), Inches(2.6), font_size=16,
    )
    add_note(
        slide,
        "Start the results by showing that the data behave sensibly at a descriptive level. "
        "Point out the diagonal pattern in the heat map: targets map to old, lures map to similar, and foils map to new. "
        "This is a credibility check for all later analyses, because meaningful REC and LDI scores depend on correct use of the response categories."
    )

    # Slide 9
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Results: Encoding RT", "Boundary position clearly slowed encoding when task rules shifted", owner="Renu")
    fit_image(slide, FIGURES / "phase1_encoding_rt.png", Inches(0.65), Inches(1.55), Inches(7.9), Inches(5.35))
    add_metric_card(
        slide, Inches(8.75), Inches(1.75), Inches(3.8), Inches(1.45),
        "Item + Task Shift",
        [
            f"ANOVA: F(2,96) = {float(both_rt['f_value']):.2f}, p < .001",
            f"Means: post {fmt3(encoding_map[('both','post')]['mean'])}, mid {fmt3(encoding_map[('both','mid')]['mean'])}, pre {fmt3(encoding_map[('both','pre')]['mean'])}",
        ],
        accent=TEAL,
    )
    add_metric_card(
        slide, Inches(8.75), Inches(3.4), Inches(3.8), Inches(1.45),
        "Task Shift Only",
        [
            f"ANOVA: F(2,104) = {float(task_rt['f_value']):.2f}, p < .001",
            f"Means: post {fmt3(encoding_map[('task_only','post')]['mean'])}, mid {fmt3(encoding_map[('task_only','mid')]['mean'])}, pre {fmt3(encoding_map[('task_only','pre')]['mean'])}",
        ],
        accent=ORANGE,
    )
    add_metric_card(
        slide, Inches(8.75), Inches(5.05), Inches(3.8), Inches(1.45),
        "Item Shift Only",
        [
            f"ANOVA: F(2,110) = {float(item_rt['f_value']):.2f}, p = {float(item_rt['p_value']):.3f}",
            "No reliable post-boundary slowing in this condition.",
        ],
        accent=BLUE,
    )
    add_note(
        slide,
        "This is the strongest result in the report. "
        "Explain that when the task rule changed, post-boundary items took longer to encode than mid or pre items. "
        "The effect is strong in both the combined item-plus-task condition and the task-only condition, but not in item-only. "
        "Interpretation: boundaries appear to increase processing demands primarily when the task context changes."
    )

    both_rec = row_for("both", "REC")
    task_rec = row_for("task_only", "REC")
    both_ldi = row_for("both", "LDI")
    task_ldi = row_for("task_only", "LDI")
    item_ldi = row_for("item_only", "LDI")

    # Slide 10
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Results: Recognition and LDI", "Recognition showed a post-boundary cost, but LDI did not show the expected pre-boundary benefit", owner="Renu")
    fit_image(slide, FIGURES / "phase1_memory_indices.png", Inches(0.55), Inches(1.55), Inches(7.25), Inches(5.45))
    add_metric_card(
        slide, Inches(8.0), Inches(1.7), Inches(4.45), Inches(2.0),
        "Recognition memory (REC)",
        [
            f"`both`: F(2,96) = {float(both_rec['f_value']):.2f}, p = {float(both_rec['p_value']):.3f}",
            f"Post REC in `both`: {fmt3(memory_map[('both','REC','post')]['mean'])} vs mid {fmt3(memory_map[('both','REC','mid')]['mean'])} and pre {fmt3(memory_map[('both','REC','pre')]['mean'])}",
            f"`task_only`: F(2,104) = {float(task_rec['f_value']):.2f}, p = {float(task_rec['p_value']):.3f}",
        ],
        accent=BLUE,
    )
    add_metric_card(
        slide, Inches(8.0), Inches(4.0), Inches(4.45), Inches(2.0),
        "Lure discrimination (LDI)",
        [
            f"`item_only`: F(2,110) = {float(item_ldi['f_value']):.2f}, p = {float(item_ldi['p_value']):.3f}",
            f"`both`: F(2,96) = {float(both_ldi['f_value']):.2f}, p = {float(both_ldi['p_value']):.3f}",
            f"`task_only`: F(2,104) = {float(task_ldi['f_value']):.2f}, p = {float(task_ldi['p_value']):.3f}",
        ],
        accent=ORANGE,
    )
    add_note(
        slide,
        "Separate the two outcomes clearly. "
        "For REC, highlight the post-boundary cost, especially in the item-plus-task condition where post is lower than both mid and pre. "
        "For LDI, state plainly that the expected pre-boundary advantage did not emerge reliably in any condition. "
        "That makes the study a partial replication: recognition shows the predicted direction more clearly than lure discrimination."
    )

    # Slide 11
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Interpretation of Findings", "What the exploratory pattern suggests", owner="Renu")
    fit_image(slide, FIGURES / "phase1_lure_bins.png", Inches(0.62), Inches(1.7), Inches(7.65), Inches(5.0))
    add_bullets(
        slide,
        [
            "The boundary manipulation worked most strongly when the task rule changed, which supports an event-segmentation style processing cost.",
            "The clearest memory effect was weaker recognition for post-boundary items rather than better discrimination for pre-boundary items.",
            "Lure-bin curves still behaved as expected: discrimination improved as lures became less similar, which supports overall task validity.",
            "So the dataset looks usable, but the strongest next-phase target is the post-boundary recognition effect, not a simple LDI replication.",
        ],
        Inches(8.55), Inches(1.9), Inches(3.6), Inches(4.5), font_size=15,
    )
    add_note(
        slide,
        "Use this slide to interpret rather than repeat test statistics. "
        "The main message is that the experiment produces a strong boundary-related processing effect and a moderate recognition cost after boundaries, but not the clean pre-boundary LDI advantage reported in the target paper. "
        "The lure-bin pattern is important because it shows that the MST itself is functioning normally even though one predicted directional effect was weak."
    )

    # Slide 12
    slide = prs.slides.add_slide(title_layout)
    add_background(slide)
    add_title(slide, "Conclusion and Future Plan", "Phase 1 establishes a clean starting point for deeper testing", owner="Pavan")
    add_metric_card(
        slide, Inches(0.8), Inches(1.75), Inches(3.8), Inches(3.8),
        "Conclusion",
        [
            "The data support a strong boundary-related processing effect.",
            "Recognition memory shows a post-boundary cost, strongest in `both`.",
            "The predicted pre-boundary LDI benefit did not replicate clearly.",
        ],
        accent=BLUE,
    )
    add_metric_card(
        slide, Inches(4.8), Inches(1.75), Inches(3.8), Inches(3.8),
        "Why this still matters",
        [
            "The response profiles and lure-bin curves indicate valid MST behavior.",
            "The exploratory phase narrowed the space of plausible next-step analyses.",
            "The strongest signal is now clearly identified for follow-up work.",
        ],
        accent=TEAL,
    )
    add_metric_card(
        slide, Inches(8.8), Inches(1.75), Inches(3.8), Inches(3.8),
        "Future plan",
        [
            "Test the post-boundary recognition effect more directly in Phase 2.",
            "Examine subject-level heterogeneity and possible boundary-implementation differences.",
            "Refine the analysis to understand why LDI did not show the expected directional advantage.",
        ],
        accent=GREEN,
    )
    close_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.0), Inches(0.45))
    p = close_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Thank you"
    r.font.name = "Aptos Display"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = NAVY
    add_note(
        slide,
        "Close by summarizing the study in three lines: boundaries clearly affected encoding, recognition showed the most convincing memory effect, and LDI did not cleanly replicate the expected pre-boundary benefit. "
        "Then state the forward-looking plan: Phase 2 should prioritize direct testing of the post-boundary recognition effect and investigate why the LDI pattern was weak. "
        "End by framing Phase 1 as successful exploratory groundwork rather than as a failed replication."
    )

    prs.save(OUTPUT_FILE)


if __name__ == "__main__":
    build_presentation()
